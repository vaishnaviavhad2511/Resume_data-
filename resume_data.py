import os
import re
import PyPDF2
import docx2txt
import win32com.client 
import docx
from docx import Document
from openpyxl import Workbook


name_pattern = r'\b(?:[A-Za-z]+(?:\s+[A-Za-z]+)?|"[A-Za-z]+\s+[A-Z][a-z]+")\b'
email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
phone_pattern = r'\b(?:\+91\s?)?(?:\d{10}|\d{3}-\d{3}-\d{4}|\d{5}\s\d{5})\b'

def extract_info_from_pdf(pdf_path):
    try:
        with open(pdf_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()

        name_match = re.search(name_pattern, text)
        email_match = re.search(email_pattern, text)
        phone_match = re.search(phone_pattern, text)

        name = name_match.group(0).strip() if name_match else None
        email = email_match.group(0).strip() if email_match else None
        phone = phone_match.group(0).strip() if phone_match else None

        return name, email, phone
        
    except Exception as e:
        print(f"An error occurred while extracting info from PDF: {e}")
        return None, None, None

def extract_info_from_docx(docx_path):
    try:
        doc = docx.Document(docx_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"

        name_match = re.search(name_pattern, text)
        email_match = re.search(email_pattern, text)
        phone_match = re.search(phone_pattern, text)

        name = name_match.group(0).strip() if name_match else None
        email = email_match.group(0).strip() if email_match else None
        phone = phone_match.group(0).strip() if phone_match else None

        return name, email, phone
        
    except Exception as e:
        print(f"An error occurred while extracting info from .docx file: {e}")
        return None, None, None

def extract_info_from_doc(doc_path):
    try:
        word_app = win32com.client.Dispatch("Word.Application")
        doc = word_app.Documents.Open(doc_path)
        text = ""
        for paragraph in doc.Paragraphs:
            text += paragraph.Range.Text + "\n"
        doc.Close()
        word_app.Quit()
            
        # Search for patterns in the text
        name_match = re.search(name_pattern, text)
        email_match = re.search(email_pattern, text)
        phone_match = re.search(phone_pattern, text)

        # Extract matched groups or set to None if not found
        name = name_match.group(0).strip() if name_match else None
        email = email_match.group(0).strip() if email_match else None
        phone = phone_match.group(0).strip() if phone_match else None

        return name, email, phone
        
    except Exception as e:
        print(f"An error occurred while extracting info from .doc file: {e}")
        return None, None, None
    
def update_excel_sheet(output_sheet, info):
    try:
        row = [info.get("Name", ""), info.get("Email", ""), info.get("Phone", "")]
        output_sheet.append(row)
    except Exception as e:
        print(f"Error occurred while updating Excel sheet: {e}")

def main():
    folder = r"C:/Users/user8/OneDrive/Desktop/Sample2"
    output_file = os.path.join(folder, "Resume_Data.xls")
    output_wb = Workbook()
    output_sheet = output_wb.active
    output_sheet.title = "Resume Data"

    # Write header row
    output_sheet['A1'] = 'Name'
    output_sheet['B1'] = 'Email'
    output_sheet['C1'] = 'Phone'

    for filename in os.listdir(folder):
        if filename.endswith('.pdf'):
            pdf_file = os.path.join(folder, filename)
            name, email, phone = extract_info_from_pdf(pdf_file)
            update_excel_sheet(output_sheet, {"Name": name, "Email": email, "Phone": phone})
        elif filename.endswith('.docx'):
            docx_file = os.path.join(folder, filename)
            name, email, phone = extract_info_from_docx(docx_file)
            update_excel_sheet(output_sheet, {"Name": name, "Email": email, "Phone": phone})
        elif filename.endswith('.doc'):
            doc_file = os.path.join(folder, filename)
            name, email, phone = extract_info_from_doc(doc_file)
            update_excel_sheet(output_sheet, {"Name": name, "Email": email, "Phone": phone})
    
    output_wb.save(output_file)

if __name__ == "__main__":
    main()
